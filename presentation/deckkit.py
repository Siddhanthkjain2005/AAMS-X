"""Shared PPTX helpers for the AAMS-X SIH decks.

The SIH template is the only permitted container, so nothing here creates a
presentation from scratch: every deck opens the uploaded file, keeps its master,
footer band, logo and slide geometry, and only replaces the prompt text boxes.
The uploaded template itself is never written to.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parent / "SIH2026-IDEA-Presentation-Format.pptx"
FIG = HERE / "figures"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

# Ink-on-white palette, identical to figstyle.L so slide text and chart ink agree.
INK = RGBColor(0x0D, 0x1B, 0x2A)
INK_DIM = RGBColor(0x3D, 0x51, 0x63)
MUTED = RGBColor(0x6B, 0x7D, 0x8F)
FAINT = RGBColor(0x93, 0xA3, 0xB3)
ACCENT = RGBColor(0x0F, 0x8B, 0x80)
MAG = RGBColor(0x6D, 0x4F, 0xD8)
SIGNAL = RGBColor(0x1F, 0x7F, 0xC4)
WARN = RGBColor(0xB7, 0x79, 0x1F)
BAD = RGBColor(0xC8, 0x38, 0x3C)
GOOD = RGBColor(0x2F, 0x8B, 0x52)
SIH = RGBColor(0x00, 0x70, 0xC0)
SIH_DARK = RGBColor(0x00, 0x4E, 0x86)
CARD = RGBColor(0xF4, 0xF7, 0xFA)
CARD2 = RGBColor(0xEA, 0xF0, 0xF6)
LINE = RGBColor(0xD3, 0xDD, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Arial"


def shape_by_name(slide, name: str):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def drop(shape) -> None:
    shape._element.getparent().remove(shape._element)


def drop_slide(prs: Presentation, index: int) -> None:
    """Remove a slide by index, unlinking the rel so the package stays valid."""
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rid = slides[index].get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    prs.part.drop_rel(rid)
    xml_slides.remove(slides[index])


def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box, tf


def para(tf, text, *, size=11, bold=False, color=INK, first=False, space_before=0,
         space_after=0, align=PP_ALIGN.LEFT, line=0.92, italic=False, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = color
    return p


def rich(tf, chunks, *, first=False, size=11, space_before=0, space_after=0,
         line=0.92, align=PP_ALIGN.LEFT, bullet_indent=None):
    """One paragraph made of (text, {overrides}) chunks — for inline emphasis."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    if bullet_indent is not None:
        p.level = 0
        pPr = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
        pPr.set("indent", str(-int(Inches(bullet_indent))))
        pPr.set("marL", str(int(Inches(bullet_indent))))
    for text, over in chunks:
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = over.get("font", FONT)
        f.size = Pt(over.get("size", size))
        f.bold = over.get("bold", False)
        f.italic = over.get("italic", False)
        f.color.rgb = over.get("color", INK)
    return p


def rounded(slide, x, y, w, h, *, fill=CARD, edge=LINE, lw=0.75, adj=0.06):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    sh.adjustments[0] = adj
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if edge is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = edge
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def bar(slide, x, y, w, h, *, fill=SIH):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def picture(slide, name, x, y, *, w=None, h=None):
    """Place figures/<name>.png at (x, y), scaled by whichever edge is given."""
    path = FIG / f"{name}.png"
    if w is not None:
        return slide.shapes.add_picture(str(path), Emu(int(x)), Emu(int(y)), width=Emu(int(w)))
    return slide.shapes.add_picture(str(path), Emu(int(x)), Emu(int(y)), height=Emu(int(h)))


def fit(name, *, w=None, h=None):
    """Return (w, h) in EMU for a figure scaled to one given edge."""
    from PIL import Image
    px, py = Image.open(FIG / f"{name}.png").size
    ar = px / py
    if w is not None:
        return int(w), int(w / ar)
    return int(h * ar), int(h)


def caption(slide, x, y, w, text, *, size=8, color=MUTED, align=PP_ALIGN.LEFT,
            italic=True, h=None):
    _, tf = textbox(slide, x, y, w, h if h is not None else Inches(0.22))
    para(tf, text, size=size, color=color, italic=italic, align=align, first=True, line=0.95)
